from unittest.mock import patch

from openpyxl import Workbook

from academic_planner.importing import (
    NewCurriculumPdfParser, OldCurriculumExcelParser, validate_curriculum,
)
from academic_planner.importing.new_pdf import term_to_placement
from academic_planner.policies.graduation_role import graduation_role_for


def test_credit_band_graduation_role_rule():
    assert [graduation_role_for(ects) for ects in (1, 2, 3)] == [
        "OPTIONAL", "OPTIONAL", "OPTIONAL",
    ]
    assert graduation_role_for(4) == "OPTIONAL"
    assert graduation_role_for(5) == "MANDATORY"
    assert graduation_role_for(6) == "MANDATORY"
    assert graduation_role_for(4, is_fyp=True) == "MANDATORY"
    assert graduation_role_for(4, course_type="TRAINING") == "MANDATORY"
    assert graduation_role_for(
        8, degree_credit_eligible=False
    ) == "NEEDS_REVIEW"


def test_old_xlsx_extracts_courses_terms_training_and_prerequisites(tmp_path):
    path = tmp_path / "CCE Courses Division old.xlsx"
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Curriculum"
    sheet.append(["Course Code", "Course Name", "ECTS", "Academic Year", "Semester", "Pre-requisite"])
    sheet.append(["ELCM101", "Circuits", 6, 1, "Fall", "-"])
    sheet.append(["ELCM202", "Signals", "5 ECTS", 2, "Spring", "ELCM101"])
    sheet.append(["ELTR301", "Training I", 2, 3, "Spring", "ELCM202"])
    sheet.append(["ELCM401", "Final Project I", 5, 4, "Fall", "ELCM202 or ELCM101"])
    workbook.save(path)

    parsed = OldCurriculumExcelParser().parse(path)
    assert [course.code for course in parsed.courses] == ["ELCM101", "ELCM202", "ELTR301", "ELCM401"]
    assert parsed.courses[1].ects == 5
    assert parsed.courses[2].reference_semester == "TRAINING_1"
    assert parsed.courses[3].is_fyp
    assert parsed.prerequisite_rules[-1].logic == "OR"
    assert parsed.prerequisite_rules[-1].prerequisite_codes == ("ELCM202", "ELCM101")
    assert parsed.degree_requirements == [
        "ELCM101", "ELCM202", "ELTR301", "ELCM401",
    ]
    assert validate_curriculum(parsed).publishable


def test_old_xlsx_preserves_ambiguous_prerequisite_text(tmp_path):
    path = tmp_path / "old.xlsx"
    workbook = Workbook(); sheet = workbook.active
    sheet.append(["Code", "Title", "ECTS", "Year", "Semester", "Pre-requisite"])
    sheet.append(["ELCC300", "Networks", 5, 3, "Fall", "Department approval or equivalent electronics course"])
    workbook.save(path)
    parsed = OldCurriculumExcelParser().parse(path)
    rule = parsed.prerequisite_rules[0]
    assert rule.ambiguous and not rule.prerequisite_codes
    assert rule.source_text == "Department approval or equivalent electronics course"
    assert any(issue.code == "AMBIGUOUS_PREREQUISITE" for issue in validate_curriculum(parsed).warnings)


def test_old_xlsx_real_layout_supports_terms_crdts_four_digit_codes_and_language_options(tmp_path):
    path = tmp_path / "CCE Courses Division old.xlsx"
    workbook = Workbook(); sheet = workbook.active
    sheet.append(["Term", "Course Code", "Course Name", "Pre-requisite", "Crdts", None])
    sheet.append([3, "CLEN/CLFR321", "Advanced English/French 1", None, 3, 3])
    sheet.append([9, "ELCM900", "FYP 1", None, 8, 8])
    sheet.append([10, "ELCM1007", "Introduction to Marketing", "ELCM900", 3, 3])
    sheet.append([None, None, "Total Credits", None, None, 14])
    workbook.save(path)

    parsed = OldCurriculumExcelParser().parse(path)
    assert [(course.code, course.language_track) for course in parsed.courses[:2]] == [
        ("CLEN321", "EN"), ("CLFR321", "FR")]
    final = next(course for course in parsed.courses if course.code == "ELCM1007")
    assert (final.reference_year, final.reference_semester) == (5, "SPRING")
    assert parsed.total_ects == parsed.stated_total_ects == 14
    assert parsed.degree_requirements == ["ELCM900"]


def test_term_conversion_is_exact():
    assert term_to_placement(1) == (1, "FALL")
    assert term_to_placement(2) == (1, "SPRING")
    assert term_to_placement(9) == (5, "FALL")
    assert term_to_placement(10) == (5, "SPRING")


def test_new_pdf_extracts_term_nine_course_and_explicit_prerequisite(tmp_path):
    path = tmp_path / "BE_CCE-per_sem-Fall2026.pdf"; path.write_bytes(b"fake-pdf")
    class Page:
        def extract_text(self, **_kwargs):
            return "Term IX\nCCE640 Image Processing 5 Pr(CCE443)\nCCE690 Final Year Project II 10 Pr(CCE640)"
    class Reader:
        def __init__(self, _path): self.pages = [Page()]
    with patch("academic_planner.importing.new_pdf.PdfReader", Reader):
        parsed = NewCurriculumPdfParser().parse(path)
    image = next(course for course in parsed.courses if course.code == "CCE640")
    assert (image.reference_year, image.reference_semester, image.ects) == (5, "FALL", 5)
    assert parsed.prerequisite_rules[0].prerequisite_codes == ("CCE443",)
    report = validate_curriculum(parsed)
    assert {issue.code for issue in report.fatal_errors} == {"BROKEN_PREREQUISITE"}


def test_new_pdf_splits_paired_columns_and_joins_wrapped_prerequisite(tmp_path):
    path = tmp_path / "BE_CCE-per_sem-Fall2026.pdf"; path.write_bytes(b"fake-pdf")
    def paired(left, right=""): return left.ljust(76) + right
    layout = "\n".join([
        paired("Term I, 5cr", "Term II, 5cr"),
        paired("MAT 230-Introductory Calculus, 5cr", "MAT 231-Calculus II, 5cr, Pr(MAT 230)"),
        paired("Term VII, 5cr", "Term VIII, 3cr"),
        paired("CCE 552-Microprocessor, 5cr, Pr(CCE", "CCE 562-Embedded Systems, 3cr, Pr(CCE552)"),
        paired("441)"),
    ])
    class Page:
        def extract_text(self, **_kwargs): return layout
    class Reader:
        def __init__(self, _path): self.pages = [Page()]
    with patch("academic_planner.importing.new_pdf.PdfReader", Reader):
        parsed = NewCurriculumPdfParser().parse(path)
    assert [course.code for course in parsed.courses] == ["MAT230", "CCE552", "MAT231", "CCE562"]
    placements = {course.code: (course.reference_year, course.reference_semester) for course in parsed.courses}
    assert placements == {"MAT230": (1, "FALL"), "MAT231": (1, "SPRING"),
                          "CCE552": (4, "FALL"), "CCE562": (4, "SPRING")}
    rules = {rule.course_code: rule.prerequisite_codes for rule in parsed.prerequisite_rules}
    assert rules == {"MAT231": ("MAT230",), "CCE552": ("CCE441",), "CCE562": ("CCE552",)}
    assert parsed.degree_requirements == ["MAT230", "CCE552", "MAT231"]


def test_duplicate_and_cycle_are_fatal(tmp_path):
    path = tmp_path / "old.xlsx"
    workbook = Workbook(); sheet = workbook.active
    sheet.append(["Code", "Title", "ECTS", "Year", "Semester", "Pre-requisite"])
    sheet.append(["ELCM100", "One", 5, 1, "Fall", "ELCM200"])
    sheet.append(["ELCM200", "Two", 5, 1, "Spring", "ELCM100"])
    sheet.append(["ELCM200", "Duplicate", 5, 2, "Fall", "-"])
    workbook.save(path)
    codes = {issue.code for issue in validate_curriculum(OldCurriculumExcelParser().parse(path)).fatal_errors}
    assert {"DUPLICATE_COURSE", "PREREQUISITE_CYCLE"} <= codes
